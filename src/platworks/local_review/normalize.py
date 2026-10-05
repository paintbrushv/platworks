"""Documented local layouts, explicit mappings, and validation; no financial math."""

import csv
import io
import re
import zipfile
from datetime import date
from pathlib import PurePath

from .common import MAX_FILE, decode, digest, finding, refuse, text

CATEGORIES = (
    "rental income",
    "concessions",
    "bad debt",
    "other income",
    "payroll",
    "repairs & maintenance",
    "utilities",
    "marketing",
    "administrative",
    "taxes",
    "insurance",
    "management fees",
)
POLICY_KEYS = {
    "data_class",
    "preparer",
    "policy_id",
    "policy_version",
    "policy_date",
    "policy_note",
    "unit_use_note",
    "reconciliation_note",
}
DELTAS = {
    "rent_growth_delta_bps",
    "exit_cap_delta_bps",
    "vacancy_delta_bps",
    "opex_growth_delta_bps",
}
ACQ_KEYS = POLICY_KEYS | {"expected_units", "down_units", "downside"}
OPS_KEYS = POLICY_KEYS | {
    "account_mapping",
    "mapping_note",
    "parent_report",
    "correction_reason",
    "property",
    "period",
    "unit_count",
}
GL_HEADERS = ("property", "period", "account_code", "account_name", "category", "amount")


def sources(files):
    if not isinstance(files, dict) or not 1 <= len(files) <= 4:
        refuse("INVALID_INPUT", "Select one to four supported source files.")
    result = {}
    for role, item in files.items():
        if role not in {"inputs", "dataset", "actuals", "budgets", "snapshot"}:
            refuse("UNSUPPORTED_LAYOUT", "Unknown source role.")
        name, raw = item
        name = text(name, "filename", limit=200)
        if (
            PurePath(name).name != name
            or "/" in name
            or "\\" in name
            or any(ord(char) < 32 or ord(char) == 127 for char in name)
        ):
            refuse("INVALID_INPUT", "Use filenames without directory components.")
        if not isinstance(raw, bytes) or not 0 < len(raw) <= MAX_FILE:
            refuse("INPUT_LIMIT", "Each source must contain at most 2 MiB.")
        suffix = PurePath(name).suffix.lower()
        if suffix not in {".json", ".csv", ".xlsx"}:
            refuse("UNSUPPORTED_LAYOUT", "Use canonical JSON or documented CSV/XLSX layouts.")
        result[role] = {
            "filename": name,
            "sha256": digest(raw),
            "bytes": len(raw),
            "format": suffix[1:],
        }
    return result


def table(raw, extension, headers, sheet_name):
    if extension == "csv":
        try:
            rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig")), strict=True))
        except (UnicodeError, csv.Error):
            refuse("UNSUPPORTED_LAYOUT", "Use a UTF-8 CSV with the documented header.")
    elif extension == "xlsx":
        from openpyxl import load_workbook

        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if (
                    len(archive.infolist()) > 2000
                    or sum(i.file_size for i in archive.infolist()) > 16 * 1024 * 1024
                ):
                    refuse("INPUT_LIMIT", "Expanded workbook exceeds 16 MiB.")
            workbook = load_workbook(
                io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            try:
                if workbook.sheetnames != [sheet_name]:
                    refuse("UNSUPPORTED_LAYOUT", f"Use one worksheet named {sheet_name}.")
                sheet = workbook[sheet_name]
                # Vendor writers can underreport the stored dimensions. Trust
                # the actual bounded row stream, never silently truncate money.
                sheet.reset_dimensions()
                rows = []
                for row in sheet.iter_rows():
                    if len(rows) >= 10001 or len(row) != len(headers):
                        refuse("INPUT_LIMIT", "Workbook dimensions exceed the documented layout.")
                    if any(cell.data_type == "f" for cell in row):
                        refuse("FORMULA_CELL", "Replace formulas with reviewed source values.")
                    rows.append([cell.value if cell.value is not None else "" for cell in row])
            finally:
                workbook.close()
        except (OSError, ValueError, KeyError, zipfile.BadZipFile):
            refuse("UNSUPPORTED_LAYOUT", "Cannot read the documented XLSX layout.")
    else:
        refuse("UNSUPPORTED_LAYOUT", "This role requires a CSV or XLSX table.")
    if not rows or tuple(rows[0]) != headers:
        refuse("UNSUPPORTED_LAYOUT", "The source header must match the documented layout exactly.")
    if len(rows) > 10001:
        refuse("INPUT_LIMIT", "At most 10,000 source rows are supported.")
    for row in rows[1:]:
        if len(row) != len(headers) or any(not isinstance(value, str) for value in row):
            refuse("AMBIGUOUS_CELL", "Use text cells, including decimal-string money and codes.")
    return [dict(zip(headers, row, strict=True)) for row in rows[1:]]


def leaves(value, pointer=""):
    if isinstance(value, dict) and value:
        for key, item in value.items():
            yield from leaves(item, pointer + "/" + key.replace("~", "~0").replace("/", "~1"))
    elif isinstance(value, list) and value:
        for i, item in enumerate(value):
            yield from leaves(item, f"{pointer}/{i}")
    else:
        yield pointer, value


def acquisition_input(raw, extension):
    if extension == "json":
        data = decode(raw)
        return data, [
            {"target": p, "locator": {"json_pointer": p}, "value": v} for p, v in leaves(data)
        ]
    rows = table(raw, extension, ("pointer", "value_json"), "Inputs")
    tree, mappings, seen = {}, [], set()
    for number, row in enumerate(rows, 2):
        pointer = row["pointer"]
        if not pointer.startswith("/") or len(pointer) > 1000 or re.search(r"~(?![01])", pointer):
            refuse("INVALID_POINTER", "Use nonempty JSON pointers with ~0/~1 escaping.")
        parts = tuple(p.replace("~1", "/").replace("~0", "~") for p in pointer[1:].split("/"))
        if any(parts[:i] in seen for i in range(1, len(parts) + 1)) or any(
            previous[: len(parts)] == parts for previous in seen
        ):
            refuse(
                "OVERLAPPING_MAPPING",
                "Duplicate or overlapping source mappings require correction.",
            )
        seen.add(parts)
        value = decode(row["value_json"])
        if isinstance(value, (dict, list)) and value:
            refuse(
                "UNSUPPORTED_LAYOUT", "Use one scalar per pointer; empty containers are allowed."
            )
        target = tree
        for part in parts[:-1]:
            target = target.setdefault(part, {})
        target[parts[-1]] = (value,)
        mappings.append(
            {
                "target": pointer,
                "locator": {"row": number, **({"sheet": "Inputs"} if extension == "xlsx" else {})},
                "value": value,
            }
        )

    def materialize(node):
        if isinstance(node, tuple):
            return node[0]
        if node and all(k.isdecimal() for k in node):
            if set(node) != {str(i) for i in range(len(node))}:
                refuse("INVALID_POINTER", "Array indexes must start at zero with no gaps.")
            return [materialize(node[str(i)]) for i in range(len(node))]
        return {k: materialize(v) for k, v in node.items()}

    return materialize(tree), mappings


def normalize(kind, files, settings):
    refs = sources(files)
    if kind not in {"acquisition", "operations"} or not isinstance(settings, dict):
        refuse("INVALID_INPUT", "Choose acquisition or operations and provide review settings.")
    allowed = ACQ_KEYS if kind == "acquisition" else OPS_KEYS
    if set(settings) - allowed or not POLICY_KEYS <= set(settings):
        refuse("INVALID_INPUT", "Review settings do not match the local review contract.")
    blockers, warnings = [], []
    for field in sorted(POLICY_KEYS):
        value = settings[field]
        if not isinstance(value, str) or len(value) > 2000:
            refuse("INVALID_INPUT", "Review text fields must be at most 2,000 characters.")
        if not value.strip():
            code = "UNIT_USE_EVIDENCE" if field == "unit_use_note" else "MISSING_EVIDENCE"
            blockers.append(finding(code, f"Provide {field.replace('_', ' ')}.", path=field))
    if settings["data_class"] not in {"synthetic", "user_supplied"}:
        refuse("INVALID_INPUT", "Declare synthetic or user_supplied source data.")
    try:
        if date.fromisoformat(settings["policy_date"]).isoformat() != settings["policy_date"]:
            raise ValueError()
    except ValueError:
        blockers.append(finding("POLICY_DATE", "Provide a YYYY-MM-DD policy as-of date."))
    if kind == "acquisition":
        if set(files) != {"inputs"} or set(settings) != ACQ_KEYS:
            refuse(
                "UNSUPPORTED_LAYOUT",
                "An acquisition requires one canonical input and explicit settings.",
            )
        data, mappings = acquisition_input(files["inputs"][1], refs["inputs"]["format"])
        for mapping in mappings:
            mapping.update(source="inputs", source_sha256=refs["inputs"]["sha256"])
        if not isinstance(data, dict):
            refuse("UNSUPPORTED_LAYOUT", "Acquisition inputs must form a JSON object.")
        if (
            not isinstance(data.get("metadata"), dict)
            or not isinstance(data.get("unit_cohorts"), list)
            or any(not isinstance(row, dict) for row in data["unit_cohorts"])
        ):
            refuse("INVALID_INPUT", "Provide canonical metadata and a list of unit cohorts.")
        grid = data.get("time_grid")
        if not isinstance(grid, dict) or any(
            not isinstance(grid.get(key), str)
            or not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", grid[key])
            for key in ("analysis_start_date", "analysis_end_date")
        ):
            refuse("INVALID_INPUT", "Provide a canonical YYYY-MM analysis time grid.")
        first, last = (grid[key] for key in ("analysis_start_date", "analysis_end_date"))
        months = (int(last[:4]) - int(first[:4])) * 12 + int(last[5:]) - int(first[5:]) + 1
        if not 1 <= months <= 600 or len(data["unit_cohorts"]) > 1000:
            refuse(
                "INPUT_LIMIT",
                "Local acquisition review supports 1–600 months and up to 1,000 cohorts.",
            )
        from engine.validator import validate_deal

        validation = validate_deal(data).to_dict()
        for item in validation["issues"]:
            (blockers if item["severity"] == "ERROR" else warnings).append(
                finding(item["code"], item["message"], path=item["path"])
            )
        explicit = {
            "purchase_assumptions": (
                "purchase_price",
                "closing_costs",
                "equity_contribution",
                "total_equity_basis",
            ),
            "debt_terms": ("commitment", "rate", "amort_years", "io_months"),
            "exit_assumptions": ("exit_cap_rate", "sale_cost_percent", "exit_month"),
            "growth_assumptions": ("growth_type", "annual_growth_rate"),
            "fund_assumptions": (
                "sponsor_equity_pct",
                "lp_equity_pct",
                "preferred_return",
                "acquisition_fee_pct",
                "asset_management_fee_pct",
                "disposition_fee_pct",
                "annual_partnership_expenses",
                "partnership_closing_costs",
            ),
        }
        for section, fields in explicit.items():
            value = data.get(section)
            if not isinstance(value, dict) or any(value.get(k) is None for k in fields):
                blockers.append(
                    finding(
                        "EXPLICIT_ASSUMPTIONS",
                        "Supply explicit "
                        + section
                        + " fields; this reviewed path does not infer financial defaults.",
                        path="/" + section,
                    )
                )
        growth = data.get("growth_assumptions")
        if not isinstance(growth, dict) or growth.get("growth_type") != "annual_compound":
            blockers.append(
                finding(
                    "GROWTH_POLICY",
                    "The initial downside workflow supports "
                    "an explicit annual_compound growth policy.",
                )
            )
        for section in ("opex_table", "fund_assumptions"):
            if not data.get(section):
                blockers.append(
                    finding(
                        "EXPLICIT_ASSUMPTIONS",
                        "Provide explicit " + section + " assumptions.",
                        path="/" + section,
                    )
                )
        expenses = data.get("opex_table")
        if isinstance(expenses, list) and any(
            not isinstance(row, dict) or row.get("growth_rate") is None for row in expenses
        ):
            blockers.append(
                finding(
                    "EXPLICIT_ASSUMPTIONS",
                    "Supply an explicit growth_rate for every expense "
                    "before applying downside deltas.",
                    path="/opex_table",
                )
            )
        warnings.append(
            finding(
                "OPTIONAL_SCHEDULES",
                "Inspect optional capital, renovation and "
                "revenue schedules in the effective inputs. Omitted schedules remain "
                "absent. The installed engine's defaults still apply to omitted optional "
                "fields (including optional loan terms); the review adds no source evidence.",
            )
        )
        subject = text(data.get("metadata", {}).get("deal_id"), "canonical deal_id", limit=120)
        units = settings["expected_units"]
        if type(units) is not int or not 1 <= units <= 100000:
            refuse("INVALID_INPUT", "Provide a positive integer expected unit count.")
        cohort_units = [c.get("unit_count") for c in data.get("unit_cohorts", [])]
        if any(type(n) is not int for n in cohort_units) or sum(cohort_units) != units:
            blockers.append(
                finding("UNIT_RECONCILIATION", "Cohort counts must match declared units.")
            )
        if type(settings["down_units"]) is not int or settings["down_units"] != 0:
            blockers.append(
                finding(
                    "DOWN_UNITS",
                    "This initial acquisition workflow requires explicit "
                    "zero down units; unresolved or nonzero down schedules are unsupported.",
                )
            )
        delta = settings["downside"]
        if (
            not isinstance(delta, dict)
            or set(delta) != DELTAS
            or any(type(v) is not int or abs(v) > 2000 for v in delta.values())
        ):
            refuse(
                "INVALID_POLICY",
                "Provide all four downside deltas as integer basis points (±2000).",
            )
        if (
            delta["rent_growth_delta_bps"] > 0
            or any(delta[k] < 0 for k in DELTAS - {"rent_growth_delta_bps"})
            or not any(delta.values())
        ):
            refuse(
                "INVALID_POLICY",
                "Downside requires lower rent growth or higher cap, vacancy or costs.",
            )
        return {
            "subject": subject,
            "period": None,
            "normalized": data,
            "mappings": mappings,
            "sources": refs,
            "blockers": blockers,
            "warnings": warnings,
        }
    for field in ("mapping_note", "correction_reason"):
        if not isinstance(settings.get(field, ""), str) or len(settings.get(field, "")) > 2000:
            refuse(
                "INVALID_INPUT", "Mapping and correction notes must be text under 2,001 characters."
            )
    if settings.get("parent_report") is not None and (
        not isinstance(settings["parent_report"], str)
        or not re.fullmatch(r"[0-9a-f]{64}", settings["parent_report"])
    ):
        refuse("INVALID_INPUT", "Select a retained parent report or leave the selection empty.")
    if "dataset" in files:
        if set(files) != {"dataset"} or refs["dataset"]["format"] != "json":
            refuse(
                "UNSUPPORTED_LAYOUT", "Use one operations dataset JSON or GL files with a snapshot."
            )
        data = decode(files["dataset"][1])
        mappings = [
            {
                "target": p,
                "locator": {"json_pointer": p},
                "value": v,
                "source": "dataset",
                "source_sha256": refs["dataset"]["sha256"],
            }
            for p, v in leaves(data)
        ]
    else:
        if set(files) != {"actuals", "budgets", "snapshot"} or refs["snapshot"]["format"] != "json":
            refuse("UNSUPPORTED_LAYOUT", "Select actual and budget GL tables plus occupancy JSON.")
        subject = text(settings.get("property"), "property", limit=120)
        period = text(settings.get("period"), "period", limit=7)
        data = {
            "property": subject,
            "period": period,
            "unit_count": settings.get("unit_count"),
            "currency": "USD",
            "expense_convention": "positive_costs",
            "snapshot": decode(files["snapshot"][1]),
        }
        mappings = []
        for role in ("actuals", "budgets"):
            rows = table(files[role][1], refs[role]["format"], GL_HEADERS, "GL")
            data[role] = []
            for number, row in enumerate(rows, 2):
                if row["property"] != subject or row["period"] != period:
                    blockers.append(
                        finding(
                            "PERIOD_CONFLICT",
                            "Every GL row must match this property "
                            "and period; split overlapping exports first.",
                        )
                    )
                line = {k: row[k] for k in ("account_code", "account_name", "category", "amount")}
                data[role].append(line)
                mappings.append(
                    {
                        "target": f"/{role}/{number - 2}",
                        "source": role,
                        "source_sha256": refs[role]["sha256"],
                        "value": line.copy(),
                        "locator": {
                            "row": number,
                            **({"sheet": "GL"} if refs[role]["format"] == "xlsx" else {}),
                        },
                    }
                )
        mappings.extend(
            {
                "target": "/snapshot" + p,
                "source": "snapshot",
                "value": v,
                "source_sha256": refs["snapshot"]["sha256"],
                "locator": {"json_pointer": p},
            }
            for p, v in leaves(data["snapshot"])
        )
    required = {
        "property",
        "period",
        "currency",
        "expense_convention",
        "unit_count",
        "actuals",
        "budgets",
        "snapshot",
    }
    if not isinstance(data, dict) or set(data) != required:
        refuse(
            "UNSUPPORTED_LAYOUT", "Operations dataset fields must match the documented contract."
        )
    subject = text(data["property"], "property", limit=120)
    if not isinstance(data["period"], str) or not re.fullmatch(
        r"\d{4}-(0[1-9]|1[0-2])", data["period"]
    ):
        refuse("INVALID_INPUT", "Use one YYYY-MM operating period.")
    mapping = settings.get("account_mapping", {})
    if not isinstance(mapping, dict) or any(v not in CATEGORIES for v in mapping.values()):
        refuse("INVALID_MAPPING", "Map accounts to supported operating categories.")
    codes, changed = set(), False
    for role in ("actuals", "budgets"):
        if not isinstance(data[role], list) or len(data[role]) > 10000:
            refuse("INVALID_INPUT", "Each GL requires a list of at most 10,000 rows.")
        for row in data[role]:
            if not isinstance(row, dict) or set(row) != {
                "account_code",
                "account_name",
                "category",
                "amount",
            }:
                refuse("INVALID_INPUT", "Invalid GL row shape.")
            code = text(row["account_code"], "account code", limit=64)
            codes.add(code)
            if code in mapping and mapping[code] != row["category"]:
                row["category"] = mapping[code]
                changed = True
            if row["category"] not in CATEGORIES:
                blockers.append(
                    finding("UNMAPPED_ACCOUNT", f"Choose a category for account {code}.")
                )
    if set(mapping) - codes:
        refuse("INVALID_MAPPING", "Mapping names an account absent from the source files.")
    if changed and not settings.get("mapping_note", "").strip():
        blockers.append(finding("MAPPING_EVIDENCE", "Explain changed account mappings."))
    if data["snapshot"] is None:
        blockers.append(finding("MISSING_SNAPSHOT", "Provide a reconciled occupancy snapshot."))
    return {
        "subject": subject,
        "period": data["period"],
        "normalized": data,
        "mappings": mappings,
        "sources": refs,
        "blockers": blockers,
        "warnings": warnings,
    }
