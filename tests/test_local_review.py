"""Real producers, synthetic files, and explicit simulated human decisions."""

import copy
import json
from importlib.resources import files

import pytest

from platworks.local_review.common import ReviewError
from platworks.local_review.store import Workspace


def acquisition():
    raw = files("platworks").joinpath("data/synthetic_acquisition.json").read_bytes()
    settings = {
        "data_class": "synthetic",
        "preparer": "Synthetic preparer",
        "policy_id": "example-policy",
        "policy_version": "1",
        "policy_date": "2026-10-05",
        "policy_note": "Synthetic assumptions for workflow testing",
        "expected_units": 100,
        "down_units": 0,
        "unit_use_note": "Synthetic roll: 100 residential units; no down or non-revenue units",
        "reconciliation_note": "Reviewed each canonical field against this fabricated source",
        "downside": {
            "rent_growth_delta_bps": -50,
            "exit_cap_delta_bps": 25,
            "vacancy_delta_bps": 200,
            "opex_growth_delta_bps": 25,
        },
    }
    return {"inputs": ("synthetic.json", raw)}, settings


def operations(amount="0.30", **settings):
    data = {
        "property": "synthetic-ops",
        "period": "2026-05",
        "currency": "USD",
        "expense_convention": "positive_costs",
        "unit_count": 10,
        "actuals": [
            {
                "account_code": "4000",
                "account_name": "Rent",
                "category": "rental income",
                "amount": amount,
            }
        ],
        "budgets": [
            {
                "account_code": "4000",
                "account_name": "Rent",
                "category": "rental income",
                "amount": "0.29",
            }
        ],
        "snapshot": {
            "as_of_date": "2026-05-31",
            "occupied_units": 9,
            "vacant_units": 1,
            "down_units": 0,
        },
    }
    return {"dataset": ("synthetic.json", json.dumps(data).encode())}, {
        "data_class": "synthetic",
        "preparer": "Synthetic preparer",
        "policy_id": "monthly-close",
        "policy_version": "1",
        "policy_date": "2026-10-05",
        "policy_note": "USD; positive expense costs",
        "unit_use_note": "All 10 units reconciled to the fabricated occupancy record",
        "reconciliation_note": "Complete matched actual and budget accounts for the month",
        "account_mapping": {},
        "mapping_note": "",
        "parent_report": None,
        "correction_reason": "",
        **settings,
    }


def decide(workspace, draft, action):
    return workspace.decide(
        action,
        draft["id"],
        draft["review_sha256"],
        reviewer="Synthetic independent reviewer",
        note="I reviewed sources, mappings, assumptions, and every exception.",
        confirmed=True,
    )


def test_acquisition_requires_two_reviews_and_freezes_original(tmp_path):
    workspace = Workspace(tmp_path / "review")
    source, settings = acquisition()
    draft = workspace.prepare("acquisition", source, settings)
    assert draft["status"] == "awaiting_execution_review"
    assert draft["result"] is None
    assert not draft["blockers"]
    with pytest.raises(ReviewError, match="EXECUTION_REQUIRED"):
        decide(workspace, draft, "issue")
    result = decide(workspace, draft, "execute")
    assert result["result"]["base"]["metrics"]
    assert result["result"]["downside"]["metrics"]
    assert result["review_sha256"] != draft["review_sha256"]
    with pytest.raises(ReviewError, match="STALE_REVIEW"):
        decide(workspace, draft, "issue")
    issued = decide(workspace, result, "issue")
    report = workspace.report(issued["report_id"])
    assert report["certified"] is False
    assert report["kind"] == "acquisition"
    assert len(report["approvals"]) == 2
    assert report["result"]["effective_inputs"]["base"] == json.loads(source["inputs"][1])
    assert Workspace(tmp_path / "review").report(issued["report_id"]) == report
    assert decide(workspace, result, "issue")["report_id"] == issued["report_id"]


def test_ops_correction_keeps_original_and_uses_rust_cent_bridge(tmp_path):
    workspace = Workspace(tmp_path / "review")
    first = workspace.prepare("operations", *operations())
    assert not first["blockers"]
    assert first["result"]["variance"]["noi_bridge"]["noi_variance"] == "0.01"
    issued = decide(workspace, first, "issue")
    original = workspace.report(issued["report_id"])
    correction = workspace.prepare(
        "operations",
        *operations(
            "0.31", parent_report=issued["report_id"], correction_reason="Correct source row"
        ),
    )
    assert correction["result"]["changes"]["actual_noi"] == "0.01"
    revised = decide(workspace, correction, "issue")
    assert workspace.report(revised["report_id"])["parent_report"] == issued["report_id"]
    assert workspace.report(issued["report_id"]) == original
    assert len(workspace.list_reports()) == 2


def test_changed_mapping_invalidates_prior_review(tmp_path):
    workspace = Workspace(tmp_path / "review")
    original = workspace.prepare("operations", *operations())
    changed = workspace.prepare(
        "operations",
        *operations(
            account_mapping={"4000": "other income"}, mapping_note="Reviewed reclassification"
        ),
    )
    assert changed["id"] != original["id"]
    with pytest.raises(ReviewError, match="STALE_DRAFT"):
        decide(workspace, original, "issue")


def test_restoring_old_inputs_requires_a_new_review_generation(tmp_path):
    workspace = Workspace(tmp_path / "review")
    original = workspace.prepare("operations", *operations())
    workspace.prepare("operations", *operations("0.31"))
    restored = workspace.prepare("operations", *operations())
    assert restored["id"] != original["id"]
    assert restored["active"] is True
    assert workspace.prepare("operations", *operations())["id"] == restored["id"]
    with pytest.raises(ReviewError, match="STALE_DRAFT"):
        decide(workspace, original, "issue")
    assert decide(workspace, restored, "issue")["report_id"]


def test_existing_unrelated_database_is_not_adopted(tmp_path):
    import sqlite3

    root = tmp_path / "review"
    root.mkdir()
    with sqlite3.connect(root / "review.sqlite") as connection:
        connection.execute("CREATE TABLE original_private_data (value TEXT)")
    before = (root / "review.sqlite").read_bytes()
    with pytest.raises(ReviewError, match="WORKSPACE_VERSION"):
        Workspace(root)
    assert (root / "review.sqlite").read_bytes() == before


def test_missing_evidence_and_down_units_block_execution(tmp_path):
    source, settings = acquisition()
    settings.update(unit_use_note="", down_units=2)
    draft = Workspace(tmp_path / "review").prepare("acquisition", source, settings)
    assert {x["code"] for x in draft["blockers"]} >= {"UNIT_USE_EVIDENCE", "DOWN_UNITS"}


def test_overlapping_acquisition_periods_block(tmp_path):
    source, settings = acquisition()
    data = json.loads(source["inputs"][1])
    data["physical_vacancy_curve"].append(copy.deepcopy(data["physical_vacancy_curve"][0]))
    data["physical_vacancy_curve"][-1]["vacancy_rate"] = 0.06
    source["inputs"] = ("overlap.json", json.dumps(data).encode())
    draft = Workspace(tmp_path / "review").prepare("acquisition", source, settings)
    assert any("SEGMENT_VACANCY" in x["code"] for x in draft["blockers"])


def test_missing_budget_never_becomes_zero(tmp_path):
    source, settings = operations()
    data = json.loads(source["dataset"][1])
    data["budgets"] = []
    source["dataset"] = ("incomplete.json", json.dumps(data).encode())
    draft = Workspace(tmp_path / "review").prepare("operations", source, settings)
    assert draft["blockers"]
    with pytest.raises(ReviewError, match="REVIEW_BLOCKED"):
        decide(Workspace(tmp_path / "review"), draft, "issue")


def test_false_confirmation_cannot_approve(tmp_path):
    workspace = Workspace(tmp_path / "review")
    draft = workspace.prepare("operations", *operations())
    for reviewer, confirmed in [("Independent", False), ("Synthetic preparer", False)]:
        with pytest.raises(ReviewError, match="HUMAN_REVIEW_REQUIRED"):
            workspace.decide(
                "issue",
                draft["id"],
                draft["review_sha256"],
                reviewer=reviewer,
                note="Reviewed",
                confirmed=confirmed,
            )


def test_duplicate_json_keys_and_unsupported_layout_refuse(tmp_path):
    workspace = Workspace(tmp_path / "review")
    _, settings = acquisition()
    for name, raw in [
        ("duplicate.json", b'{"schema_version":"0.1","schema_version":"bad"}'),
        ("broker.pdf", b"unsupported"),
    ]:
        with pytest.raises(ReviewError):
            workspace.prepare("acquisition", {"inputs": (name, raw)}, settings)


def test_changed_producer_refuses_before_financial_action(tmp_path, monkeypatch):
    from platworks.local_review import producers

    workspace = Workspace(tmp_path / "review")
    draft = workspace.prepare("operations", *operations())
    changed = copy.deepcopy(draft["producer_identity"])
    changed["boxscore-exact"]["binary_sha256"] = "0" * 64
    monkeypatch.setattr(producers, "producer_identity", lambda: changed)
    assert workspace.get(draft["id"])["producer_current"] is False
    with pytest.raises(ReviewError, match="STALE_PRODUCER"):
        decide(workspace, draft, "issue")
    assert workspace.list_reports() == []


def test_crash_before_commit_can_resume_without_double_issue(tmp_path, monkeypatch):
    from platworks.local_review import producers

    workspace = Workspace(tmp_path / "review")
    draft = workspace.prepare("operations", *operations())
    real_issue = producers.operations_issue

    def interrupted(*args):
        real_issue(*args)
        raise RuntimeError("simulated interruption before workspace commit")

    monkeypatch.setattr(producers, "operations_issue", interrupted)
    with pytest.raises(RuntimeError):
        decide(workspace, draft, "issue")
    assert Workspace(tmp_path / "review").list_reports() == []
    monkeypatch.setattr(producers, "operations_issue", real_issue)
    issued = decide(workspace, draft, "issue")
    assert decide(workspace, draft, "issue")["report_id"] == issued["report_id"]
    assert len(workspace.list_reports()) == 1


def test_correction_must_select_current_report(tmp_path):
    workspace = Workspace(tmp_path / "review")
    first = decide(workspace, workspace.prepare("operations", *operations()), "issue")
    second = workspace.prepare(
        "operations",
        *operations("0.31", parent_report=first["report_id"], correction_reason="First correction"),
    )
    decide(workspace, second, "issue")
    stale = workspace.prepare(
        "operations",
        *operations("0.32", parent_report=first["report_id"], correction_reason="Stale correction"),
    )
    assert any(b["code"] == "STALE_PARENT" for b in stale["blockers"])
    with pytest.raises(ReviewError, match="REVIEW_BLOCKED"):
        decide(workspace, stale, "issue")
    assert len(workspace.list_reports()) == 2


def test_source_blob_tampering_is_detected(tmp_path):
    import sqlite3

    workspace = Workspace(tmp_path / "review")
    draft = workspace.prepare("operations", *operations())
    with sqlite3.connect(workspace.database) as connection:
        # Deliberately bypass the normal immutable-table guard to test read integrity.
        connection.execute("DROP TRIGGER blobs_no_UPDATE")
        connection.execute(
            "UPDATE blobs SET body=? WHERE sha=?",
            (b"altered source", draft["sources"]["dataset"]["sha256"]),
        )
    with pytest.raises(ReviewError, match="INTEGRITY_ERROR"):
        workspace.get(draft["id"])


def test_csv_and_xlsx_acquisition_mappings_preserve_all_values():
    import csv
    import io

    from openpyxl import Workbook

    from platworks.local_review.normalize import acquisition_input, leaves

    source, _ = acquisition()
    data = json.loads(source["inputs"][1])
    rows = [(p, json.dumps(value)) for p, value in leaves(data)]
    csv_data = io.StringIO()
    writer = csv.writer(csv_data)
    writer.writerow(["pointer", "value_json"])
    writer.writerows(rows)
    decoded, mappings = acquisition_input(csv_data.getvalue().encode(), "csv")
    assert decoded == data
    assert mappings[0]["locator"] == {"row": 2}
    book = Workbook()
    book.active.title = "Inputs"
    book.active.append(["pointer", "value_json"])
    for row in rows:
        book.active.append(row)
    raw = io.BytesIO()
    book.save(raw)
    decoded, mappings = acquisition_input(raw.getvalue(), "xlsx")
    assert decoded == data
    assert mappings[0]["locator"] == {"row": 2, "sheet": "Inputs"}


def test_excel_formula_is_an_explicit_blocker():
    import io

    from openpyxl import Workbook

    from platworks.local_review.normalize import acquisition_input

    book = Workbook()
    book.active.title = "Inputs"
    book.active.append(["pointer", "value_json"])
    book.active.append(["/purchase_assumptions/purchase_price", "=2+2"])
    raw = io.BytesIO()
    book.save(raw)
    with pytest.raises(ReviewError, match="FORMULA_CELL"):
        acquisition_input(raw.getvalue(), "xlsx")


def test_operations_table_period_conflicts_block(tmp_path):
    source, settings = operations()
    snapshot = json.loads(source["dataset"][1])["snapshot"]
    header = "property,period,account_code,account_name,category,amount\n"
    files = {
        "actuals": (
            "actuals.csv",
            (header + "synthetic-ops,2026-06,4000,Rent,rental income,0.30\n").encode(),
        ),
        "budgets": (
            "budgets.csv",
            (header + "synthetic-ops,2026-05,4000,Rent,rental income,0.29\n").encode(),
        ),
        "snapshot": ("occupancy.json", json.dumps(snapshot).encode()),
    }
    settings.update(property="synthetic-ops", period="2026-05", unit_count=10)
    draft = Workspace(tmp_path / "review").prepare("operations", files, settings)
    assert any(b["code"] == "PERIOD_CONFLICT" for b in draft["blockers"])


@pytest.mark.parametrize(
    "section,value",
    [
        ("growth_assumptions", []),
        ("metadata", []),
        ("unit_cohorts", [None]),
    ],
)
def test_malformed_acquisition_sections_refuse_cleanly(tmp_path, section, value):
    source, settings = acquisition()
    data = json.loads(source["inputs"][1])
    data[section] = value
    source["inputs"] = ("invalid.json", json.dumps(data).encode())
    workspace = Workspace(tmp_path / "review")
    try:
        draft = workspace.prepare("acquisition", source, settings)
    except ReviewError:
        return
    assert draft["blockers"]
    with pytest.raises(ReviewError, match="REVIEW_BLOCKED"):
        decide(workspace, draft, "execute")


def test_missing_financial_assumptions_cannot_use_engine_defaults(tmp_path):
    source, settings = acquisition()
    data = json.loads(source["inputs"][1])
    del data["purchase_assumptions"]["closing_costs"]
    source["inputs"] = ("missing.json", json.dumps(data).encode())
    workspace = Workspace(tmp_path / "review")
    draft = workspace.prepare("acquisition", source, settings)
    assert any(b["code"] == "EXPLICIT_ASSUMPTIONS" for b in draft["blockers"])
    with pytest.raises(ReviewError, match="REVIEW_BLOCKED"):
        decide(workspace, draft, "execute")


def test_concurrent_reviewers_issue_exactly_once(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier

    workspace = Workspace(tmp_path / "review")
    draft = workspace.prepare("operations", *operations())
    barrier = Barrier(2)

    def review(reviewer):
        barrier.wait(timeout=5)
        return workspace.decide(
            "issue",
            draft["id"],
            draft["review_sha256"],
            reviewer=reviewer,
            note="Synthetic concurrency test",
            confirmed=True,
        )

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(review, ("Synthetic reviewer A", "Synthetic reviewer B")))
    assert results[0]["report_id"] == results[1]["report_id"]
    assert len(workspace.list_reports()) == 1
    assert len(workspace.report(results[0]["report_id"])["approvals"]) == 1


@pytest.mark.parametrize("amount", ["0.001", "92233720368547758.08", "NaN"])
def test_operating_money_precision_and_overflow_block_issue(tmp_path, amount):
    workspace = Workspace(tmp_path / "review")
    draft = workspace.prepare("operations", *operations(amount))
    assert draft["blockers"]
    with pytest.raises(ReviewError, match="REVIEW_BLOCKED"):
        decide(workspace, draft, "issue")


def test_excessive_horizon_refuses_before_unbounded_validation(tmp_path):
    source, settings = acquisition()
    data = json.loads(source["inputs"][1])
    data["time_grid"]["analysis_end_date"] = "2090-06"
    source["inputs"] = ("long-horizon.json", json.dumps(data).encode())
    with pytest.raises(ReviewError, match="INPUT_LIMIT"):
        Workspace(tmp_path / "review").prepare("acquisition", source, settings)


@pytest.mark.parametrize(
    "section,field",
    [
        ("fund_assumptions", "preferred_return"),
        ("opex_table", "growth_rate"),
    ],
)
def test_missing_scenario_or_fund_assumption_is_blocked(tmp_path, section, field):
    source, settings = acquisition()
    data = json.loads(source["inputs"][1])
    target = data[section][0] if section == "opex_table" else data[section]
    del target[field]
    source["inputs"] = ("missing-policy.json", json.dumps(data).encode())
    draft = Workspace(tmp_path / "review").prepare("acquisition", source, settings)
    assert any(b["code"] == "EXPLICIT_ASSUMPTIONS" for b in draft["blockers"])


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_operations_tables_can_issue_and_preserve_source_locations(tmp_path, extension):
    import csv
    import io

    from openpyxl import Workbook

    from platworks.local_review.normalize import GL_HEADERS

    source, settings = operations()
    data = json.loads(source["dataset"][1])
    selected = {"snapshot": ("occupancy.json", json.dumps(data["snapshot"]).encode())}
    for role in ("actuals", "budgets"):
        rows = [list(GL_HEADERS)] + [
            [
                data["property"],
                data["period"],
                row["account_code"],
                row["account_name"],
                row["category"],
                row["amount"],
            ]
            for row in data[role]
        ]
        if extension == "csv":
            stream = io.StringIO()
            csv.writer(stream).writerows(rows)
            raw = stream.getvalue().encode()
        else:
            book = Workbook()
            book.active.title = "GL"
            for row in rows:
                book.active.append(row)
            stream = io.BytesIO()
            book.save(stream)
            raw = stream.getvalue()
        selected[role] = (role + "." + extension, raw)
    settings.update(property=data["property"], period=data["period"], unit_count=data["unit_count"])
    workspace = Workspace(tmp_path / "review")
    draft = workspace.prepare("operations", selected, settings)
    assert not draft["blockers"]
    report = workspace.report(decide(workspace, draft, "issue")["report_id"])
    assert report["result"]["variance"]["noi_bridge"]["noi_variance"] == "0.01"
    assert report["mappings"][0]["locator"]["row"] == 2
    if extension == "xlsx":
        assert report["mappings"][0]["locator"]["sheet"] == "GL"
    for role, (_, raw) in selected.items():
        assert workspace.source(report["sources"][role]["sha256"])[1] == raw


def test_declared_user_supplied_path_keeps_separate_contract(tmp_path):
    # Fabricated bytes exercise the declaration branch; this is not a real-file pilot.
    source, settings = operations(data_class="user_supplied")
    workspace = Workspace(tmp_path / "review")
    draft = workspace.prepare("operations", source, settings)
    report = workspace.report(decide(workspace, draft, "issue")["report_id"])
    assert report["contract_version"] == "local-reviewed-draft/1.0.0"
    assert report["data_class"] == "user_supplied"
    assert report["certified"] is False


def test_policy_edit_cannot_replace_frozen_original_thesis(tmp_path):
    workspace = Workspace(tmp_path / "review")
    source, settings = acquisition()
    first = workspace.prepare("acquisition", source, settings)
    calculated = decide(workspace, first, "execute")
    issued = decide(workspace, calculated, "issue")
    original = workspace.report(issued["report_id"])
    settings["policy_note"] = "Synthetic updated acquisition view"
    changed = workspace.prepare("acquisition", source, settings)
    assert changed["id"] != first["id"]
    assert changed["result"] is None
    updated = decide(workspace, changed, "execute")
    with pytest.raises(ReviewError, match="ORIGINAL_THESIS_EXISTS"):
        decide(workspace, updated, "issue")
    assert workspace.report(issued["report_id"]) == original


def test_running_workspace_requires_restart_after_code_changes(tmp_path, monkeypatch):
    from platworks.local_review import producers

    workspace = Workspace(tmp_path / "review")
    workspace.prepare("operations", *operations())
    changed = copy.deepcopy(producers.producer_identity())
    changed["platworks"]["content_sha256"] = "0" * 64
    monkeypatch.setattr(producers, "producer_identity", lambda: changed)
    with pytest.raises(ReviewError, match="STALE_PRODUCER"):
        workspace.prepare("operations", *operations("0.31"))


def test_xlsx_underreported_dimensions_cannot_hide_gl_rows(tmp_path):
    import io
    import re
    import zipfile

    from openpyxl import Workbook

    from platworks.local_review.normalize import GL_HEADERS

    book = Workbook()
    book.active.title = "GL"
    book.active.append(list(GL_HEADERS))
    for amount in ("0.10", "0.20"):
        book.active.append(["synthetic-ops", "2026-05", "4000", "Rent", "rental income", amount])
    saved = io.BytesIO()
    book.save(saved)
    rewritten = io.BytesIO()
    with zipfile.ZipFile(saved) as source, zipfile.ZipFile(rewritten, "w") as target:
        for entry in source.infolist():
            body = source.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                body = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:F2"', body)
            target.writestr(entry, body)
    source, settings = operations()
    snapshot = json.loads(source["dataset"][1])["snapshot"]
    selected = {
        "actuals": ("actuals.xlsx", rewritten.getvalue()),
        "budgets": (
            "budgets.csv",
            (
                ",".join(GL_HEADERS) + "\nsynthetic-ops,2026-05,4000,Rent,rental income,0.29\n"
            ).encode(),
        ),
        "snapshot": ("snapshot.json", json.dumps(snapshot).encode()),
    }
    settings.update(property="synthetic-ops", period="2026-05", unit_count=10)
    draft = Workspace(tmp_path / "review").prepare("operations", selected, settings)
    assert not draft["blockers"]
    assert draft["result"]["variance"]["noi_bridge"]["noi_variance"] == "0.01"
    assert len(draft["normalized"]["actuals"]) == 2


def test_history_pages_keep_older_reports_reachable(tmp_path):
    workspace = Workspace(tmp_path / "review")
    first = decide(workspace, workspace.prepare("operations", *operations()), "issue")
    revised = workspace.prepare(
        "operations",
        *operations(
            "0.31", parent_report=first["report_id"], correction_reason="Synthetic correction"
        ),
    )
    second = decide(workspace, revised, "issue")
    assert workspace.list_reports(offset=0, limit=1)[0]["id"] == second["report_id"]
    older = workspace.list_reports(offset=1, limit=1)
    assert older[0]["id"] == first["report_id"]
    assert (
        workspace.report(older[0]["id"])["result"]["variance"]["noi_bridge"]["actual_noi"] == "0.30"
    )
    assert workspace.list_drafts(offset=1, limit=1)[0]["id"] == first["id"]
    assert workspace.get(first["id"])["status"] == "issued"


def test_padded_account_codes_can_be_mapped_through_review(tmp_path):
    source, settings = operations()
    data = json.loads(source["dataset"][1])
    for role in ("actuals", "budgets"):
        data[role][0].update(account_code=" 4000 ", category="unknown export label")
    raw = json.dumps(data).encode()
    source["dataset"] = ("padded.json", raw)
    workspace = Workspace(tmp_path / "review")
    draft = workspace.prepare("operations", source, settings)
    # The browser derives its mapping key from this normalized code.
    code = draft["normalized"]["actuals"][0]["account_code"]
    settings.update(account_mapping={code: "rental income"}, mapping_note="Reviewed rent account")
    corrected = workspace.prepare("operations", source, settings)
    assert not corrected["blockers"]
    assert corrected["normalized"]["actuals"][0]["account_code"] == "4000"
    report = workspace.report(decide(workspace, corrected, "issue")["report_id"])
    assert report["result"]["variance"]["noi_bridge"]["noi_variance"] == "0.01"
    assert workspace.source(report["sources"]["dataset"]["sha256"])[1] == raw


def test_csv_stops_reading_at_the_row_limit(monkeypatch):
    from platworks.local_review import normalize

    real_reader = normalize.csv.reader

    def bounded_reader(*args, **kwargs):
        for count, row in enumerate(real_reader(*args, **kwargs), 1):
            if count > 10002:
                raise AssertionError("Reader consumed records beyond the rejection boundary")
            yield row

    monkeypatch.setattr(normalize.csv, "reader", bounded_reader)
    with pytest.raises(ReviewError, match="INPUT_LIMIT"):
        normalize.table(b"header\n" + b"value\n" * 20000, "csv", ("header",), "Unused")


@pytest.mark.parametrize("kind", ["acquisition", "operations"])
def test_json_leaf_limit_refuses_before_saving_expanded_mappings(tmp_path, kind):
    source, settings = acquisition() if kind == "acquisition" else operations()
    role = "inputs" if kind == "acquisition" else "dataset"
    data = json.loads(source[role][1])
    if kind == "acquisition":
        data["metadata"]["synthetic_invalid_values"] = [0] * 10001
    else:
        data["snapshot"] = [0] * 10001
    raw = json.dumps(data).encode()
    assert len(raw) < 2 * 1024 * 1024
    source[role] = ("synthetic-oversized.json", raw)
    workspace = Workspace(tmp_path / "review")
    with pytest.raises(ReviewError, match="INPUT_LIMIT"):
        workspace.prepare(kind, source, settings)
    assert workspace.list_drafts() == []


@pytest.mark.parametrize("shape", ["deep", "long_locators"])
def test_json_mapping_locations_are_bounded(tmp_path, shape):
    source, settings = operations()
    data = json.loads(source["dataset"][1])
    if shape == "deep":
        value = 0
        for _ in range(66):
            value = {"nested": value}
    else:
        value = {"x" * 900: [0] * 2000}
    data["snapshot"] = value
    source["dataset"] = ("synthetic-shape.json", json.dumps(data).encode())
    with pytest.raises(ReviewError, match="INPUT_LIMIT"):
        Workspace(tmp_path / "review").prepare("operations", source, settings)
